import io
from datetime import datetime
import streamlit as st
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go

# ---------------------------------------------------------
# Safe Imports with Friendly Cloud Deployment Fallbacks
# ---------------------------------------------------------
try:
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    REPORTLAB_AVAILABLE = True
except ModuleNotFoundError:
    REPORTLAB_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    DOCX_AVAILABLE = True
except ModuleNotFoundError:
    DOCX_AVAILABLE = False

# ---------------------------------------------------------
# 1. Page Configuration & Professional CSS Styling
# ---------------------------------------------------------
st.set_page_config(
    page_title="High-Risk TB Analytical Intelligence Suite",
    page_icon="🩺",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&display=swap');
    
    html, body, [class*="css"] {
        font-family: 'Plus Jakarta Sans', sans-serif;
    }
    
    .stApp {
        background-color: #F8FAFC;
    }
    
    .block-container {
        padding-top: 1rem !important;
        padding-bottom: 2rem !important;
        padding-left: 2rem !important;
        padding-right: 2rem !important;
        max-width: 100% !important;
    }
    
    /* Clean Light Theme Sidebar with Deep Black Text for Perfect Visibility */
    section[data-testid="stSidebar"] {
        background-color: #F1F5F9 !important;
        border-right: 1px solid #CBD5E1 !important;
    }
    section[data-testid="stSidebar"] * {
        color: #0F172A !important;
    }
    section[data-testid="stSidebar"] label {
        font-size: 0.82rem !important;
        font-weight: 700 !important;
        color: #0F172A !important;
        margin-bottom: 3px !important;
    }
    section[data-testid="stSidebar"] .stMarkdown p {
        color: #0F172A !important;
    }
    section[data-testid="stSidebar"] [data-testid="stWidgetLabel"] p {
        color: #0F172A !important;
        font-weight: 700 !important;
    }
    
    /* File Uploader in Sidebar */
    div[data-testid="stFileUploader"] section {
        background-color: #FFFFFF !important;
        border: 1.5px dashed #94A3B8 !important;
        border-radius: 8px !important;
        padding: 8px 12px !important;
    }
    div[data-testid="stFileUploader"] section small {
        display: none !important;
    }
    div[data-testid="stFileUploader"] button {
        background: #2563EB !important;
        color: #FFFFFF !important;
        border: none !important;
        border-radius: 4px !important;
        padding: 4px 10px !important;
        font-size: 0.78rem !important;
        font-weight: 600 !important;
    }
    div[data-testid="stFileUploader"] button * {
        color: #FFFFFF !important;
    }

    /* Multiselect Tags Styling */
    span[data-baseweb="tag"] {
        background-color: #2563EB !important;
        border-radius: 4px !important;
        font-size: 0.75rem !important;
    }
    span[data-baseweb="tag"] * {
        color: #FFFFFF !important;
    }

    /* Action Buttons in Sidebar */
    section[data-testid="stSidebar"] .stButton button {
        background-color: #FFFFFF !important;
        border: 1.5px solid #94A3B8 !important;
        color: #0F172A !important;
        font-weight: 700 !important;
        font-size: 0.76rem !important;
        padding: 5px 10px !important;
        border-radius: 6px !important;
        transition: all 0.2s ease;
    }
    section[data-testid="stSidebar"] .stButton button:hover {
        background-color: #E2E8F0 !important;
        border-color: #2563EB !important;
    }

    .sidebar-brand-pill {
        display: inline-block;
        background: #1D4ED8;
        color: #FFFFFF !important;
        padding: 4px 10px;
        border-radius: 6px;
        font-size: 0.72rem;
        font-weight: 800;
        letter-spacing: 0.5px;
        text-transform: uppercase;
        margin-bottom: 8px;
    }
    .sidebar-brand-pill * {
        color: #FFFFFF !important;
    }
    .sidebar-heading {
        font-size: 0.9rem;
        font-weight: 800;
        color: #0F172A !important;
        margin-top: 10px;
        margin-bottom: 6px;
        display: flex;
        align-items: center;
        gap: 6px;
    }
    .sidebar-sep {
        height: 1px;
        background-color: #CBD5E1;
        margin: 12px 0 10px 0;
    }

    .top-hero-card {
        background: #FFFFFF;
        border: 1px solid #E2E8F0;
        border-left: 6px solid #2563EB;
        border-radius: 10px;
        padding: 16px 22px;
        margin-bottom: 16px;
        box-shadow: 0 1px 3px rgba(0, 0, 0, 0.05);
        display: flex;
        justify-content: space-between;
        align-items: center;
    }
    .hero-title-main {
        font-size: 1.35rem;
        font-weight: 800;
        color: #0F172A;
        margin: 0;
        line-height: 1.2;
    }
    .hero-meta-subtitle {
        font-size: 0.82rem;
        color: #64748B;
        margin-top: 4px;
        font-weight: 500;
    }
    .live-status-pill {
        background-color: #ECFDF5;
        color: #059669;
        border: 1px solid #A7F3D0;
        padding: 4px 12px;
        border-radius: 20px;
        font-size: 0.75rem;
        font-weight: 700;
        display: inline-flex;
        align-items: center;
        gap: 6px;
    }

    .metric-card-wrapper {
        background: #FFFFFF;
        border: 1px solid #E2E8F0;
        border-radius: 10px;
        padding: 16px 18px;
        box-shadow: 0 1px 3px rgba(0, 0, 0, 0.04);
        position: relative;
        overflow: hidden;
    }
    .metric-accent-line {
        position: absolute;
        top: 0;
        left: 0;
        right: 0;
        height: 4px;
    }
    .accent-blue { background: #2563EB; }
    .accent-green { background: #10B981; }
    .accent-amber { background: #F59E0B; }
    .accent-purple { background: #8B5CF6; }
    .accent-rose { background: #E11D48; }

    .metric-label-text {
        font-size: 0.72rem;
        font-weight: 700;
        text-transform: uppercase;
        letter-spacing: 0.5px;
        color: #64748B;
        margin-bottom: 4px;
    }
    .metric-value-text {
        font-size: 1.7rem;
        font-weight: 800;
        color: #0F172A;
        line-height: 1.1;
    }
    .metric-sub-note {
        font-size: 0.72rem;
        color: #94A3B8;
        margin-top: 4px;
        font-weight: 500;
    }

    .stTabs [data-baseweb="tab-list"] {
        gap: 8px;
        border-bottom: 2px solid #E2E8F0;
        margin-bottom: 16px;
    }
    .stTabs [data-baseweb="tab"] {
        background-color: transparent;
        border-radius: 6px 6px 0 0;
        font-size: 0.88rem;
        font-weight: 600;
        color: #64748B;
        padding: 8px 16px;
    }
    .stTabs [aria-selected="true"] {
        color: #2563EB !important;
        border-bottom: 2px solid #2563EB !important;
        background-color: #EFF6FF !important;
    }

    .insight-box {
        background: #FFFFFF;
        border: 1px solid #E2E8F0;
        border-radius: 8px;
        padding: 16px;
        margin-bottom: 12px;
        box-shadow: 0 1px 2px rgba(0,0,0,0.03);
    }
    .insight-title {
        font-size: 0.95rem;
        font-weight: 700;
        color: #0F172A;
        margin-bottom: 8px;
        display: flex;
        align-items: center;
        gap: 6px;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# ---------------------------------------------------------
# 2. Ingestion & Schema Matching Helpers
# ---------------------------------------------------------
def find_column(df, possible_names):
    col_map = {str(c).strip().lower().replace("_", " ").replace("-", " "): c for c in df.columns}
    for name in possible_names:
        clean_name = name.strip().lower().replace("_", " ").replace("-", " ")
        if clean_name in col_map:
            return col_map[clean_name]
    return None

def normalize_text(series):
    return (
        series.fillna("Unknown")
        .astype(str)
        .str.strip()
        .replace({"": "Unknown", "nan": "Unknown", "None": "Unknown", "NAN": "Unknown"})
        .str.title()
    )

def safe_unique_list(series_list):
    combined = set()
    for s in series_list:
        for item in s:
            if pd.notna(item):
                item_str = str(item).strip()
                if item_str and item_str.lower() not in ["none", "nan", "null"]:
                    combined.add(item_str)
    return sorted(list(combined), key=lambda x: str(x).lower())

@st.cache_data(show_spinner=False)
def load_any_table(file_bytes, filename):
    bio = io.BytesIO(file_bytes)
    if filename.lower().endswith(".csv"):
        try:
            return pd.read_csv(bio, encoding="utf-8")
        except UnicodeDecodeError:
            bio.seek(0)
            return pd.read_csv(bio, encoding="latin1")
    else:
        return pd.read_excel(bio)

# ---------------------------------------------------------
# 3. ReportLab Enterprise PDF Generator
# ---------------------------------------------------------
def generate_pdf_report(title_text, subtitle_text, sections, landscape_mode=True):
    if not REPORTLAB_AVAILABLE:
        return b""

    buffer = io.BytesIO()
    page_size = landscape(letter) if landscape_mode else letter
    margin = 0.35 * inch

    doc = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#0F172A"),
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#64748B"),
    )
    section_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9.5,
        leading=12,
        textColor=colors.HexColor("#1E293B"),
        spaceBefore=6,
        spaceAfter=4,
    )
    body_text_style = ParagraphStyle(
        "BodyNarrative",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=10.5,
        textColor=colors.HexColor("#334155"),
        spaceAfter=4,
    )
    th_style = ParagraphStyle(
        "TableHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7,
        leading=8.5,
        textColor=colors.white,
        alignment=1,
    )
    tb_style = ParagraphStyle(
        "TableBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=6.5,
        leading=8,
        textColor=colors.HexColor("#1E293B"),
    )
    tb_bold_style = ParagraphStyle(
        "TableBodyBold",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=6.5,
        leading=8,
        textColor=colors.HexColor("#0F172A"),
    )

    story = []
    story.append(Paragraph(title_text, title_style))
    story.append(Paragraph(f"{subtitle_text}  |  Generated: {datetime.now().strftime('%d %b %Y, %I:%M %p')}", subtitle_style))
    story.append(Spacer(1, 6))

    printable_width = (doc.pagesize[0] - doc.leftMargin - doc.rightMargin)

    for sec in sections:
        sec_title = sec.get("title")
        sec_narrative = sec.get("narrative")
        df_sec = sec.get("df")
        summary_stats = sec.get("summary_stats", None)

        if sec_title:
            story.append(Paragraph(sec_title, section_style))

        if sec_narrative:
            for line in sec_narrative:
                story.append(Paragraph(f"• {line}", body_text_style))
            story.append(Spacer(1, 4))

        if summary_stats:
            lbl_row = []
            val_row = []
            for k, v in summary_stats.items():
                lbl_row.append(Paragraph(f"<b>{k.upper()}</b>", ParagraphStyle("KLbl", fontName="Helvetica-Bold", fontSize=6, textColor=colors.HexColor("#64748B"), alignment=1)))
                val_row.append(Paragraph(f"<b>{v}</b>", ParagraphStyle("KVal", fontName="Helvetica-Bold", fontSize=9.5, textColor=colors.HexColor("#1D4ED8"), alignment=1)))
            
            kpi_table = Table([lbl_row, val_row], colWidths=[printable_width / len(summary_stats)] * len(summary_stats))
            kpi_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F1F5F9")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(kpi_table)
            story.append(Spacer(1, 6))

        if df_sec is not None and not df_sec.empty:
            headers = [Paragraph(str(col), th_style) for col in df_sec.columns]
            table_rows = [headers]

            for _, row in df_sec.iterrows():
                row_cells = []
                for idx, col in enumerate(df_sec.columns):
                    val_str = str(row[col]) if pd.notna(row[col]) else "-"
                    if idx == 0:
                        row_cells.append(Paragraph(val_str, tb_bold_style))
                    else:
                        row_cells.append(Paragraph(val_str, tb_style))
                table_rows.append(row_cells)

            num_cols = len(df_sec.columns)
            col_widths = [printable_width / num_cols] * num_cols

            t = Table(table_rows, colWidths=col_widths, repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("LEFTPADDING", (0, 0), (-1, -1), 2.5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2.5),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E2E8F0")),
            ]))
            story.append(t)

        story.append(Spacer(1, 8))

    doc.build(story)
    return buffer.getvalue()

# ---------------------------------------------------------
# 4. Executive Word (.docx) Report Generator
# ---------------------------------------------------------
def generate_word_report(title_text, subtitle_text, summary_stats, insights_list, recommendations_list, data_tables):
    if not DOCX_AVAILABLE:
        return b""

    doc = Document()
    
    for sec in doc.sections:
        sec.top_margin = Inches(0.7)
        sec.bottom_margin = Inches(0.7)
        sec.left_margin = Inches(0.7)
        sec.right_margin = Inches(0.7)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run(title_text)
    title_run.font.name = "Arial"
    title_run.font.size = Pt(17)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)
    title_p.paragraph_format.space_after = Pt(2)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(f"{subtitle_text} | Report Date: {datetime.now().strftime('%d %B %Y, %I:%M %p')}")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(9.5)
    sub_run.font.color.rgb = RGBColor(100, 116, 139)
    sub_p.paragraph_format.space_after = Pt(12)

    if summary_stats:
        h1 = doc.add_heading("1. Executive Operational Metrics & Targets", level=2)
        h1.style.font.color.rgb = RGBColor(30, 41, 59)
        
        stat_table = doc.add_table(rows=2, cols=len(summary_stats))
        stat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for idx, (k, v) in enumerate(summary_stats.items()):
            cell_lbl = stat_table.cell(0, idx)
            cell_lbl.text = k.upper()
            cell_lbl.paragraphs[0].runs[0].font.size = Pt(7.5)
            cell_lbl.paragraphs[0].runs[0].font.bold = True
            cell_lbl.paragraphs[0].runs[0].font.color.rgb = RGBColor(71, 85, 105)
            shd_lbl = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
            cell_lbl._tc.get_or_add_tcPr().append(shd_lbl)
            
            cell_val = stat_table.cell(1, idx)
            cell_val.text = str(v)
            cell_val.paragraphs[0].runs[0].font.size = Pt(11)
            cell_val.paragraphs[0].runs[0].font.bold = True
            cell_val.paragraphs[0].runs[0].font.color.rgb = RGBColor(37, 99, 235)
            shd_val = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>')
            cell_val._tc.get_or_add_tcPr().append(shd_val)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    if insights_list:
        h2 = doc.add_heading("2. Diagnostic Intelligence & Performance Patterns", level=2)
        h2.style.font.color.rgb = RGBColor(30, 41, 59)
        for item in insights_list:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(item)
            r.font.size = Pt(9.5)
            r.font.name = "Arial"
            r.font.color.rgb = RGBColor(30, 41, 59)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    if recommendations_list:
        h3 = doc.add_heading("3. Strategic Action Guide & Recommended Interventions", level=2)
        h3.style.font.color.rgb = RGBColor(30, 41, 59)
        for item in recommendations_list:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(item)
            r.font.size = Pt(9.5)
            r.font.name = "Arial"
            r.font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    sec_num = 4
    for title, df in data_tables:
        if df is not None and not df.empty:
            h_tbl = doc.add_heading(f"{sec_num}. {title}", level=2)
            h_tbl.style.font.color.rgb = RGBColor(30, 41, 59)
            sec_num += 1

            tbl = doc.add_table(rows=len(df) + 1, cols=len(df.columns))
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

            for c_idx, col_name in enumerate(df.columns):
                c = tbl.cell(0, c_idx)
                c.text = str(col_name)
                c.paragraphs[0].runs[0].font.bold = True
                c.paragraphs[0].runs[0].font.size = Pt(8)
                c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E293B"/>')
                c._tc.get_or_add_tcPr().append(shd)

            for r_idx, (_, row) in enumerate(df.iterrows()):
                for c_idx, col_name in enumerate(df.columns):
                    c = tbl.cell(r_idx + 1, c_idx)
                    val = str(row[col_name]) if pd.notna(row[col_name]) else "-"
                    c.text = val
                    c.paragraphs[0].runs[0].font.size = Pt(7.5)
                    fill_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
                    c._tc.get_or_add_tcPr().append(shd)

            doc.add_paragraph().paragraph_format.space_after = Pt(10)

    out_stream = io.BytesIO()
    doc.save(out_stream)
    return out_stream.getvalue()

# ---------------------------------------------------------
# 5. Sidebar File Upload Section
# ---------------------------------------------------------
with st.sidebar:
    st.markdown('<div class="sidebar-brand-pill">TB Care Analytics</div>', unsafe_allow_html=True)
    st.markdown('<div class="sidebar-heading">📁 File Ingestion</div>', unsafe_allow_html=True)
    
    patient_file = st.file_uploader(
        "1. Patient Line List (CSV/Excel)",
        type=["csv", "xlsx", "xls"],
        help="Upload CSV or Excel containing High-Risk TB patient records.",
    )
    followup_file = st.file_uploader(
        "2. Follow-up Report (CSV/Excel)",
        type=["csv", "xlsx", "xls"],
        help="Upload CSV or Excel containing follow-up completion records.",
    )
    
    if not REPORTLAB_AVAILABLE or not DOCX_AVAILABLE:
        st.caption("⚠️ Optional: Add `reportlab` & `python-docx` to `requirements.txt` for PDF/Word exports.")

    st.markdown('<div class="sidebar-sep"></div>', unsafe_allow_html=True)

# ---------------------------------------------------------
# 6. Empty State Handler (Awaiting Files)
# ---------------------------------------------------------
if not patient_file or not followup_file:
    st.markdown(
        """
        <div class="top-hero-card">
            <div>
                <h1 class="hero-title-main">High-Risk TB Patient Follow-up Monitoring Dashboard</h1>
                <div class="hero-meta-subtitle">Automated Multi-District TB Patient Follow-up & Diagnostic Guidance Suite</div>
            </div>
            <div class="live-status-pill" style="background-color: #FEF3C7; color: #D97706; border-color: #FDE68A;">
                ● Awaiting Data Upload
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.info("👈 **Please upload both data files (.csv, .xlsx, or .xls) using the sidebar to unlock interactive analytics and reports.**")
    st.stop()

# ---------------------------------------------------------
# 7. Schema Validation & Dynamic Ingestion
# ---------------------------------------------------------
try:
    df_p_raw = load_any_table(patient_file.getvalue(), patient_file.name)
    df_f_raw = load_any_table(followup_file.getvalue(), followup_file.name)
except Exception as e:
    st.error(f"❌ Error parsing input files: {e}")
    st.stop()

col_p_dist = find_column(df_p_raw, ["Current Facility District", "District", "District Name"])
col_p_tu = find_column(df_p_raw, ["Current Facility TBU", "Current Facility TU", "TU", "Tuberculosis Unit", "TBU"])
col_p_id = find_column(df_p_raw, ["Episode ID", "Patient ID", "Nikshay ID", "ID"])

col_f_dist = find_column(df_f_raw, ["District", "Current Facility District", "District Name"])
col_f_tu = find_column(df_f_raw, ["Facility Name", "Current Facility TBU", "Current Facility TU", "TU", "TBU"])
col_f_user = find_column(df_f_raw, ["User Name", "User", "Followed Up By", "Staff Name"])
col_f_count = find_column(df_f_raw, ["Followed Up Patient Count", "Follow up count", "Count", "Followed Up Count"])
col_f_id = find_column(df_f_raw, ["Episode ID", "Patient ID", "Nikshay ID", "ID"])

if not col_p_dist or not col_p_tu:
    st.error("❌ The Patient Line List must have identifiable 'District' and 'TU/TBU' columns.")
    st.stop()

if not col_f_dist:
    st.error("❌ The Follow-up Report must have an identifiable 'District' column.")
    st.stop()

# ---------------------------------------------------------
# 8. Relational Cleaning & Normalization
# ---------------------------------------------------------
df_pt = df_p_raw.copy()
df_pt["District_Clean"] = normalize_text(df_pt[col_p_dist])
df_pt["TU_Clean"] = normalize_text(df_pt[col_p_tu])

if col_p_id:
    df_pt = df_pt.dropna(subset=[col_p_id]).drop_duplicates(subset=[col_p_id])
else:
    df_pt = df_pt.drop_duplicates()

df_fu = df_f_raw.copy()
df_fu["District_Clean"] = normalize_text(df_fu[col_f_dist])
df_fu["TU_Clean"] = normalize_text(df_fu[col_f_tu]) if col_f_tu else "Unknown"
df_fu["User_Clean"] = normalize_text(df_fu[col_f_user]) if col_f_user else "Unassigned"

if col_f_count:
    df_fu["Followup_Count_Num"] = pd.to_numeric(df_fu[col_f_count], errors="coerce").fillna(0).astype(int)
else:
    df_fu["Followup_Count_Num"] = 1

has_patient_level_fu = bool(
    col_f_id and col_p_id and set(df_fu[col_f_id].dropna()).intersection(set(df_pt[col_p_id].dropna()))
)

if has_patient_level_fu:
    fu_matched_ids = set(df_fu.dropna(subset=[col_f_id])[col_f_id])
    df_pt["Follow-up Status"] = df_pt[col_p_id].isin(fu_matched_ids).map({True: "Completed", False: "Pending"})
    if col_f_user:
        fu_user_map = df_fu.dropna(subset=[col_f_id]).drop_duplicates(subset=[col_f_id]).set_index(col_f_id)[col_f_user].to_dict()
        df_pt["Follow-up User"] = df_pt[col_p_id].map(fu_user_map).fillna("Unassigned")
else:
    df_pt["Follow-up Status"] = "Aggregated Record"
    df_pt["Follow-up User"] = "See User Tab"

# ---------------------------------------------------------
# 9. Smart Searchable Multi-Select Filter Panel
# ---------------------------------------------------------
all_districts = safe_unique_list([df_pt["District_Clean"].unique(), df_fu["District_Clean"].unique()])
total_dist_count = len(all_districts)

if "selected_districts_list" not in st.session_state:
    st.session_state["selected_districts_list"] = all_districts.copy()

def do_select_all_districts():
    st.session_state["selected_districts_list"] = all_districts.copy()

def do_clear_all_districts():
    st.session_state["selected_districts_list"] = []

with st.sidebar:
    st.markdown('<div class="sidebar-heading">🔍 Scope Filtering</div>', unsafe_allow_html=True)
    
    col_b1, col_b2 = st.columns(2)
    with col_b1:
        st.button("✅ Select All", on_click=do_select_all_districts, width="stretch")
    with col_b2:
        st.button("❌ Clear All", on_click=do_clear_all_districts, width="stretch")

    selected_districts = st.multiselect(
        "📍 Type or Select District(s):",
        options=all_districts,
        default=st.session_state["selected_districts_list"],
        key="district_selector_ms",
        placeholder="Type district name to search...",
        help="Type any letter to instantly filter and select multiple districts."
    )
    st.session_state["selected_districts_list"] = selected_districts

    active_count = len(selected_districts)
    st.markdown(
        f"<div style='font-size:0.78rem; font-weight:700; color:#1E293B; margin-top:3px; margin-bottom:5px;'>"
        f"Active: <span style='color:#2563EB;'>{active_count} of {total_dist_count}</span> Districts Selected"
        f"</div>",
        unsafe_allow_html=True,
    )

    st.markdown('<div class="sidebar-sep"></div>', unsafe_allow_html=True)

    # TU Filter
    if selected_districts:
        pt_tus = df_pt[df_pt["District_Clean"].isin(selected_districts)]["TU_Clean"].unique()
        fu_tus = df_fu[df_fu["District_Clean"].isin(selected_districts)]["TU_Clean"].unique()
        available_tus = safe_unique_list([pt_tus, fu_tus])
    else:
        available_tus = []
    tu_choice = st.selectbox("Filter Tuberculosis Unit (TU)", ["All"] + available_tus)

    # Staff Filter
    if selected_districts:
        fu_user_scoped = df_fu[df_fu["District_Clean"].isin(selected_districts)].copy()
        if tu_choice != "All":
            fu_user_scoped = fu_user_scoped[fu_user_scoped["TU_Clean"] == tu_choice]
        available_users = safe_unique_list([fu_user_scoped["User_Clean"].unique()])
    else:
        available_users = []
    user_choice = st.selectbox("Filter User / Field Staff", ["All"] + available_users)

# Apply Scoping Filters to DataFrames
if selected_districts:
    filtered_pt = df_pt[df_pt["District_Clean"].isin(selected_districts)].copy()
    filtered_fu = df_fu[df_fu["District_Clean"].isin(selected_districts)].copy()

    if tu_choice != "All":
        filtered_pt = filtered_pt[filtered_pt["TU_Clean"] == tu_choice]
        filtered_fu = filtered_fu[filtered_fu["TU_Clean"] == tu_choice]

    if user_choice != "All":
        filtered_fu = filtered_fu[filtered_fu["User_Clean"] == user_choice]
        if has_patient_level_fu:
            filtered_pt = filtered_pt[filtered_pt["Follow-up User"] == user_choice]
else:
    filtered_pt = df_pt.iloc[0:0].copy()
    filtered_fu = df_fu.iloc[0:0].copy()

# ---------------------------------------------------------
# 10. Metric Computations & Aggregations
# ---------------------------------------------------------
kpi_total_patients = len(filtered_pt)

if has_patient_level_fu:
    kpi_completed = len(filtered_pt[filtered_pt["Follow-up Status"] == "Completed"])
else:
    kpi_completed = int(filtered_fu["Followup_Count_Num"].sum()) if not filtered_fu.empty else 0

kpi_pending = max(0, kpi_total_patients - kpi_completed)
kpi_rate = round((kpi_completed / kpi_total_patients * 100), 2) if kpi_total_patients > 0 else 0.0

target_benchmark_pct = 85.0
target_needed_completions = int(kpi_total_patients * (target_benchmark_pct / 100))
target_gap = max(0, target_needed_completions - kpi_completed)

def assign_priority_tier(rate):
    if rate >= 75.0:
        return "On Track"
    elif rate >= 50.0:
        return "Moderate"
    else:
        return "Critical Attention"

# District Level Aggregation
if not filtered_pt.empty:
    d_pt = filtered_pt.groupby("District_Clean").size().reset_index(name="Total High-Risk Patients")
    if has_patient_level_fu:
        d_fu = filtered_pt[filtered_pt["Follow-up Status"] == "Completed"].groupby("District_Clean").size().reset_index(name="Follow-up Completed")
    else:
        d_fu = filtered_fu.groupby("District_Clean")["Followup_Count_Num"].sum().reset_index(name="Follow-up Completed")

    district_df = pd.merge(d_pt, d_fu, on="District_Clean", how="outer").fillna(0)
    district_df["Total High-Risk Patients"] = district_df["Total High-Risk Patients"].astype(int)
    district_df["Follow-up Completed"] = district_df["Follow-up Completed"].astype(int)
    district_df["Pending Follow-up"] = (district_df["Total High-Risk Patients"] - district_df["Follow-up Completed"]).clip(lower=0)
    district_df["Follow-up %"] = (
        (district_df["Follow-up Completed"] / district_df["Total High-Risk Patients"].replace(0, pd.NA) * 100)
        .fillna(0)
        .round(2)
    )
    district_df["Priority Tier"] = district_df["Follow-up %"].apply(assign_priority_tier)
    district_df = district_df.rename(columns={"District_Clean": "District"}).sort_values(by="Total High-Risk Patients", ascending=False)
else:
    district_df = pd.DataFrame(columns=["District", "Total High-Risk Patients", "Follow-up Completed", "Pending Follow-up", "Follow-up %", "Priority Tier"])

# TU Level Aggregation
if not filtered_pt.empty:
    t_pt = filtered_pt.groupby(["District_Clean", "TU_Clean"]).size().reset_index(name="Total High-Risk Patients")
    if has_patient_level_fu:
        t_fu = filtered_pt[filtered_pt["Follow-up Status"] == "Completed"].groupby(["District_Clean", "TU_Clean"]).size().reset_index(name="Follow-up Completed")
    else:
        t_fu = filtered_fu.groupby(["District_Clean", "TU_Clean"])["Followup_Count_Num"].sum().reset_index(name="Follow-up Completed")

    tu_df = pd.merge(t_pt, t_fu, on=["District_Clean", "TU_Clean"], how="outer").fillna(0)
    tu_df["Total High-Risk Patients"] = tu_df["Total High-Risk Patients"].astype(int)
    tu_df["Follow-up Completed"] = tu_df["Follow-up Completed"].astype(int)
    tu_df["Pending Follow-up"] = (tu_df["Total High-Risk Patients"] - tu_df["Follow-up Completed"]).clip(lower=0)
    tu_df["Follow-up %"] = (
        (tu_df["Follow-up Completed"] / tu_df["Total High-Risk Patients"].replace(0, pd.NA) * 100)
        .fillna(0)
        .round(2)
    )
    tu_df["Priority Tier"] = tu_df["Follow-up %"].apply(assign_priority_tier)
    tu_df = tu_df.rename(columns={"District_Clean": "District", "TU_Clean": "TU"}).sort_values(by="Total High-Risk Patients", ascending=False)
else:
    tu_df = pd.DataFrame(columns=["District", "TU", "Total High-Risk Patients", "Follow-up Completed", "Pending Follow-up", "Follow-up %", "Priority Tier"])

# Staff Level Aggregation
if not filtered_pt.empty:
    if has_patient_level_fu:
        user_df = (
            filtered_pt[filtered_pt["Follow-up Status"] == "Completed"]
            .groupby(["District_Clean", "TU_Clean", "Follow-up User"])
            .size()
            .reset_index(name="Patients Followed Up")
            .rename(columns={"District_Clean": "District", "TU_Clean": "TU", "Follow-up User": "User"})
        )
    else:
        user_df = (
            filtered_fu.groupby(["District_Clean", "TU_Clean", "User_Clean"])["Followup_Count_Num"]
            .sum()
            .reset_index(name="Patients Followed Up")
            .rename(columns={"District_Clean": "District", "TU_Clean": "TU", "User_Clean": "User"})
        )

    total_user_fu = user_df["Patients Followed Up"].sum()
    user_df["Follow-up Share %"] = (
        (user_df["Patients Followed Up"] / total_user_fu * 100).fillna(0).round(2) if total_user_fu > 0 else 0.0
    )
    user_df = user_df.sort_values(by="Patients Followed Up", ascending=False)
else:
    user_df = pd.DataFrame(columns=["District", "TU", "User", "Patients Followed Up", "Follow-up Share %"])

# ---------------------------------------------------------
# 11. Top Hero Header Banner
# ---------------------------------------------------------
is_single_district = len(selected_districts) == 1
is_all_districts = len(selected_districts) == total_dist_count

if not selected_districts:
    scope_dist_txt = "No District Selected (0 Scope)"
elif is_all_districts:
    scope_dist_txt = f"All {total_dist_count} Program Districts"
elif is_single_district:
    scope_dist_txt = f"District: {selected_districts[0]}"
else:
    scope_dist_txt = f"{len(selected_districts)} Selected Districts"

st.markdown(
    f"""
    <div class="top-hero-card">
        <div>
            <h1 class="hero-title-main">High-Risk TB Patient Follow-up Monitoring Dashboard</h1>
            <div class="hero-meta-subtitle">
                Active Geography: <strong>{scope_dist_txt}</strong> &nbsp;|&nbsp; TU Scope: <strong>{tu_choice}</strong> &nbsp;|&nbsp; Staff: <strong>{user_choice}</strong>
            </div>
        </div>
        <div class="live-status-pill">
            <span style="font-size: 8px;">●</span> Intelligence Active
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

if not selected_districts:
    st.warning("⚠️ **All districts are currently cleared.** Type in the search box or click **'Select All'** in the sidebar to populate data.")

# ---------------------------------------------------------
# 12. Metric KPI Cards (Always Show 0 When Cleared)
# ---------------------------------------------------------
c1, c2, c3, c4 = st.columns(4)
with c1:
    st.markdown(
        f"""
        <div class="metric-card-wrapper">
            <div class="metric-accent-line accent-blue"></div>
            <div class="metric-label-text">Total High-Risk Cohort</div>
            <div class="metric-value-text">{kpi_total_patients:,}</div>
            <div class="metric-sub-note">Verified patient records</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
with c2:
    st.markdown(
        f"""
        <div class="metric-card-wrapper">
            <div class="metric-accent-line accent-green"></div>
            <div class="metric-label-text">Follow-up Completed</div>
            <div class="metric-value-text">{kpi_completed:,}</div>
            <div class="metric-sub-note">Verified completed cases</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
with c3:
    st.markdown(
        f"""
        <div class="metric-card-wrapper">
            <div class="metric-accent-line accent-amber"></div>
            <div class="metric-label-text">Pending Interventions</div>
            <div class="metric-value-text">{kpi_pending:,}</div>
            <div class="metric-sub-note">Requiring field follow-up</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
with c4:
    accent_class = "accent-green" if kpi_rate >= 75 else ("accent-amber" if kpi_rate >= 50 else "accent-rose")
    st.markdown(
        f"""
        <div class="metric-card-wrapper">
            <div class="metric-accent-line {accent_class}"></div>
            <div class="metric-label-text">Overall Coverage Rate</div>
            <div class="metric-value-text">{kpi_rate}%</div>
            <div class="metric-sub-note">Target: ≥85.0% (Gap: {target_gap:,})</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

st.write("")

# ---------------------------------------------------------
# 13. Data Analyst Diagnostic Engine
# ---------------------------------------------------------
def generate_analyst_intelligence():
    insights = []
    recommendations = []

    if not selected_districts or kpi_total_patients == 0:
        insights.append("<b>Zero District Scope:</b> No geographical units are currently selected. Type a district name or use 'Select All' in the sidebar.")
        recommendations.append("Select one or more districts from the sidebar to generate operational insights and action plans.")
        return insights, recommendations

    active_staff_count = user_df[user_df["User"] != "Unassigned"]["User"].nunique() if not user_df.empty else 0
    avg_per_staff = round(kpi_completed / active_staff_count, 1) if active_staff_count > 0 else 0

    if is_single_district:
        d_name = selected_districts[0]
        insights.append(f"<b>District Performance Profile:</b> <b>{d_name}</b> has registered a total cohort of <b>{kpi_total_patients:,}</b> high-risk patients. To date, <b>{kpi_completed:,}</b> patients ({kpi_rate}%) have verified follow-ups, leaving <b>{kpi_pending:,}</b> pending field completion.")
        
        if kpi_rate < target_benchmark_pct:
            insights.append(f"<b>Programmatic Target Gap:</b> To achieve the national benchmark of <b>{target_benchmark_pct}%</b>, {d_name} must close an additional <b>{target_gap:,}</b> pending cases.")
        else:
            insights.append(f"<b>Benchmark Status:</b> {d_name} has surpassed the {target_benchmark_pct}% national target with <b>{kpi_rate}%</b> completion. Current focus should be continuous surveillance.")

        dist_tus = tu_df[tu_df["District"] == d_name]
        if not dist_tus.empty:
            best_tu = dist_tus.sort_values(by="Follow-up %", ascending=False).iloc[0]
            lag_tu = dist_tus.sort_values(by="Follow-up %", ascending=True).iloc[0]
            
            insights.append(f"<b>Intra-District TU Disparity:</b> The leading Tuberculosis Unit is <b>{best_tu['TU']}</b> ({best_tu['Follow-up %']}% follow-up), whereas <b>{lag_tu['TU']}</b> requires critical supervisory attention at <b>{lag_tu['Follow-up %']}%</b> with <b>{lag_tu['Pending Follow-up']:,}</b> pending patients.")
            
            crit_tus = dist_tus[dist_tus["Priority Tier"] == "Critical Attention"]
            if not crit_tus.empty:
                recommendations.append(f"Immediate micro-plan deployment in <b>{len(crit_tus)} Critical Attention TU(s)</b> ({', '.join(crit_tus['TU'].head(4).tolist())}).")

        insights.append(f"<b>Field Staff Capacity:</b> Currently <b>{active_staff_count} active field workers</b> are mapped in {d_name}, recording an average workload of <b>{avg_per_staff}</b> completed visits per worker.")
        recommendations.append(f"Deploy 14-day intensive field sprints targeting the remaining <b>{kpi_pending:,} pending high-risk TB patients</b>.")
        recommendations.append("Conduct weekly physical verification of unassigned follow-up records to ensure reporting fidelity.")

    else:
        insights.append(f"<b>Macro Geographical Scope:</b> Analyzing aggregate outcomes across <b>{len(selected_districts)} districts</b> with a cumulative volume of <b>{kpi_total_patients:,} high-risk patients</b>.")
        
        if not district_df.empty:
            best_dist = district_df.sort_values(by="Follow-up %", ascending=False).iloc[0]
            lag_dist = district_df.sort_values(by="Follow-up %", ascending=True).iloc[0]
            insights.append(f"<b>Inter-District Disparity:</b> Benchmark leader is <b>{best_dist['District']}</b> ({best_dist['Follow-up %']}%), while <b>{lag_dist['District']}</b> shows lowest follow-up velocity at <b>{lag_dist['Follow-up %']}%</b> ({lag_dist['Pending Follow-up']:,} pending).")

            crit_dists = district_df[district_df["Priority Tier"] == "Critical Attention"]
            mod_dists = district_df[district_df["Priority Tier"] == "Moderate"]
            track_dists = district_df[district_df["Priority Tier"] == "On Track"]

            insights.append(f"<b>State Tier Breakdown:</b> <b>{len(track_dists)}</b> On Track (≥75%), <b>{len(mod_dists)}</b> Moderate (50-75%), and <b>{len(crit_dists)}</b> Critical Attention (<50%).")
            
            if not crit_dists.empty:
                recommendations.append(f"Issue state-level operational reviews for <b>{len(crit_dists)} Critical Attention Districts</b>: {', '.join(crit_dists['District'].tolist())}.")

            recommendations.append(f"Benchmark and replicate field outreach practices from <b>{best_dist['District']}</b> across lagging units.")
            recommendations.append("Establish cross-district staff reallocation or centralized tele-verification to balance field workloads.")

    return insights, recommendations

insights_list, recommendations_list = generate_analyst_intelligence()

clean_insights = [item.replace("<b>", "").replace("</b>", "") for item in insights_list]
clean_recs = [item.replace("<b>", "").replace("</b>", "") for item in recommendations_list]

global_kpi_summary = {
    "Scope": scope_dist_txt,
    "Cohort": f"{kpi_total_patients:,}",
    "Completed": f"{kpi_completed:,}",
    "Pending": f"{kpi_pending:,}",
    "Coverage %": f"{kpi_rate}%",
    "85% Gap": f"{target_gap:,}",
}

# ---------------------------------------------------------
# 14. Tabbed Analytics Workspace
# ---------------------------------------------------------
tab_guide, tab_dist, tab_tu, tab_staff, tab_patients = st.tabs(
    ["💡 Data Analyst Intelligence Guide", "📊 District Overview", "🏥 TU-Wise Analysis", "👥 Field Staff Performance", "📋 Patient Line-List"]
)

plotly_theme = dict(
    plot_bgcolor="#FFFFFF",
    paper_bgcolor="#FFFFFF",
    font=dict(family="Plus Jakarta Sans", size=11, color="#475569"),
    margin=dict(l=10, r=10, t=25, b=10),
)

# --- TAB 1: EXECUTIVE ANALYSIS GUIDE ---
with tab_guide:
    st.markdown("### 🩺 Public Health Executive Diagnostic Brief")
    st.caption(f"Automated Senior Data Analyst overview assessing coverage velocity, bottlenecks, and interventions for: **{scope_dist_txt}**")

    col_g1, col_g2 = st.columns([6, 6])
    with col_g1:
        st.markdown(
            f"""
            <div class="insight-box" style="border-left: 5px solid #2563EB;">
                <div class="insight-title">🔍 Diagnostic Findings & Program Observations</div>
                <div style="font-size: 0.85rem; color: #334155; line-height: 1.6;">
                    {'<br>'.join(['• ' + item for item in insights_list])}
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )
    with col_g2:
        st.markdown(
            f"""
            <div class="insight-box" style="border-left: 5px solid #10B981;">
                <div class="insight-title">🎯 Targeted Programmatic Action Guide</div>
                <div style="font-size: 0.85rem; color: #334155; line-height: 1.6;">
                    {'<br>'.join(['• ' + item for item in recommendations_list])}
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    st.markdown(f"#### 📥 Download Executive Diagnostic Guide for **{scope_dist_txt}**")
    c_gd1, c_gd2 = st.columns(2)

    if is_single_district:
        report_tables = [
            ("TU-Wise Analytical Breakdown", tu_df[tu_df["District"] == selected_districts[0]]),
            ("Staff Productivity Output", user_df[user_df["District"] == selected_districts[0]]),
        ]
        report_title = f"High-Risk TB Diagnostic Analysis: {selected_districts[0]}"
    else:
        report_tables = [
            ("District Performance Matrix", district_df),
            ("Top Tuberculosis Units (TUs)", tu_df.head(25)),
        ]
        report_title = f"High-Risk TB Comprehensive Diagnostic Analysis ({len(selected_districts)} Districts)"

    with c_gd1:
        if kpi_total_patients > 0 and REPORTLAB_AVAILABLE:
            guide_pdf_sections = [
                {
                    "title": "1. Operational KPI Scorecard",
                    "df": report_tables[0][1],
                    "summary_stats": global_kpi_summary,
                },
                {
                    "title": "2. Diagnostic Findings & Analytical Observations",
                    "narrative": clean_insights,
                },
                {
                    "title": "3. Strategic Action Guide & Recommended Interventions",
                    "narrative": clean_recs,
                },
            ]
            guide_pdf_bytes = generate_pdf_report(
                title_text=report_title,
                subtitle_text=f"Filter Scope: {scope_dist_txt} | TU: {tu_choice} | Staff: {user_choice}",
                sections=guide_pdf_sections,
            )
        else:
            guide_pdf_bytes = b""
            
        st.download_button(
            label=f"📄 Download Diagnostic Guide (PDF) - {scope_dist_txt}",
            data=guide_pdf_bytes,
            file_name=f"TB_Diagnostic_Guide_{scope_dist_txt.replace(' ', '_').replace(':', '')}.pdf",
            mime="application/pdf",
            width="stretch",
            disabled=(kpi_total_patients == 0 or not REPORTLAB_AVAILABLE),
        )

    with c_gd2:
        if kpi_total_patients > 0 and DOCX_AVAILABLE:
            guide_docx_bytes = generate_word_report(
                title_text=report_title,
                subtitle_text=f"Active Filter Scope: {scope_dist_txt} | TU: {tu_choice} | Staff: {user_choice}",
                summary_stats=global_kpi_summary,
                insights_list=clean_insights,
                recommendations_list=clean_recs,
                data_tables=report_tables,
            )
        else:
            guide_docx_bytes = b""
            
        st.download_button(
            label=f"📝 Download Diagnostic Guide (Word .docx) - {scope_dist_txt}",
            data=guide_docx_bytes,
            file_name=f"TB_Diagnostic_Guide_{scope_dist_txt.replace(' ', '_').replace(':', '')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            width="stretch",
            disabled=(kpi_total_patients == 0 or not DOCX_AVAILABLE),
        )

# --- TAB 2: DISTRICT OVERVIEW ---
with tab_dist:
    col_d1, col_d2 = st.columns([6, 6])
    with col_d1:
        st.markdown("**District Follow-up vs Pending Status**")
        fig_d = go.Figure()
        if not district_df.empty:
            fig_d.add_trace(go.Bar(name="Completed", x=district_df["District"].head(25), y=district_df["Follow-up Completed"].head(25), marker_color="#10B981"))
            fig_d.add_trace(go.Bar(name="Pending", x=district_df["District"].head(25), y=district_df["Pending Follow-up"].head(25), marker_color="#EF4444"))
        fig_d.update_layout(barmode="stack", height=320, legend=dict(orientation="h", y=1.12, x=0.25), **plotly_theme)
        fig_d.update_xaxes(showgrid=False)
        fig_d.update_yaxes(gridcolor="#F1F5F9")
        st.plotly_chart(fig_d, width="stretch")

    with col_d2:
        st.markdown("**District Performance Scorecard**")
        st.dataframe(district_df, width="stretch", height=270)
        
        c_exp_d1, c_exp_d2 = st.columns(2)
        with c_exp_d1:
            if not district_df.empty and REPORTLAB_AVAILABLE:
                dist_pdf = generate_pdf_report(
                    title_text="District-Level TB Follow-up Analysis Report",
                    subtitle_text=f"Scope: {scope_dist_txt}",
                    sections=[{"title": "District Performance Scorecard", "df": district_df, "summary_stats": global_kpi_summary}],
                )
            else:
                dist_pdf = b""
            st.download_button(
                label="📄 Download District PDF",
                data=dist_pdf,
                file_name="TB_District_Analysis_Report.pdf",
                mime="application/pdf",
                width="stretch",
                disabled=(district_df.empty or not REPORTLAB_AVAILABLE),
            )
        with c_exp_d2:
            if not district_df.empty and DOCX_AVAILABLE:
                dist_docx = generate_word_report(
                    title_text="District-Level TB Follow-up Analysis Report",
                    subtitle_text=f"Scope: {scope_dist_txt}",
                    summary_stats=global_kpi_summary,
                    insights_list=clean_insights,
                    recommendations_list=clean_recs,
                    data_tables=[("District Performance Breakdown", district_df)],
                )
            else:
                dist_docx = b""
            st.download_button(
                label="📝 Download District Word (.docx)",
                data=dist_docx,
                file_name="TB_District_Analysis.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                width="stretch",
                disabled=(district_df.empty or not DOCX_AVAILABLE),
            )

# --- TAB 3: TU ANALYSIS ---
with tab_tu:
    col_t1, col_t2 = st.columns([6, 6])
    with col_t1:
        st.markdown("**Top Tuberculosis Units (TUs) by Cohort Volume**")
        if not tu_df.empty:
            fig_t = px.bar(
                tu_df.head(20),
                x="TU",
                y=["Follow-up Completed", "Pending Follow-up"],
                barmode="group",
                color_discrete_sequence=["#2563EB", "#F59E0B"],
            )
        else:
            fig_t = go.Figure()
        fig_t.update_layout(height=320, legend=dict(orientation="h", y=1.12, x=0.25), **plotly_theme)
        fig_t.update_xaxes(showgrid=False)
        fig_t.update_yaxes(gridcolor="#F1F5F9")
        st.plotly_chart(fig_t, width="stretch")

    with col_t2:
        st.markdown("**TU Analytical Performance Table**")
        st.dataframe(tu_df, width="stretch", height=270)
        
        c_exp_t1, c_exp_t2 = st.columns(2)
        with c_exp_t1:
            if not tu_df.empty and REPORTLAB_AVAILABLE:
                tu_pdf = generate_pdf_report(
                    title_text="Tuberculosis Unit (TU) Operational Monitoring Report",
                    subtitle_text=f"District: {scope_dist_txt} | TU: {tu_choice}",
                    sections=[{"title": "TU Performance Breakdown", "df": tu_df, "summary_stats": global_kpi_summary}],
                )
            else:
                tu_pdf = b""
            st.download_button(
                label="📄 Download TU PDF",
                data=tu_pdf,
                file_name="TB_TU_Operational_Report.pdf",
                mime="application/pdf",
                width="stretch",
                disabled=(tu_df.empty or not REPORTLAB_AVAILABLE),
            )
        with c_exp_t2:
            if not tu_df.empty and DOCX_AVAILABLE:
                tu_docx = generate_word_report(
                    title_text="Tuberculosis Unit (TU) Operational Monitoring Report",
                    subtitle_text=f"District: {scope_dist_txt} | TU: {tu_choice}",
                    summary_stats=global_kpi_summary,
                    insights_list=clean_insights,
                    recommendations_list=clean_recs,
                    data_tables=[("TU Performance Table", tu_df)],
                )
            else:
                tu_docx = b""
            st.download_button(
                label="📝 Download TU Word (.docx)",
                data=tu_docx,
                file_name="TB_TU_Analysis.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                width="stretch",
                disabled=(tu_df.empty or not DOCX_AVAILABLE),
            )

# --- TAB 4: STAFF PERFORMANCE ---
with tab_staff:
    col_u1, col_u2 = st.columns([5, 7])
    with col_u1:
        st.markdown("**Staff Follow-up Output Ranking**")
        if not user_df.empty:
            fig_u = px.bar(
                user_df.head(20),
                x="Patients Followed Up",
                y="User",
                orientation="h",
                color="Patients Followed Up",
                color_continuous_scale="Blues",
            )
            fig_u.update_layout(yaxis=dict(autorange="reversed"), height=320, **plotly_theme)
        else:
            fig_u = go.Figure()
            fig_u.update_layout(height=320, **plotly_theme)
        fig_u.update_xaxes(gridcolor="#F1F5F9")
        fig_u.update_yaxes(showgrid=False)
        st.plotly_chart(fig_u, width="stretch")

    with col_u2:
        st.markdown("**Field Staff Productivity Audit**")
        st.dataframe(user_df, width="stretch", height=270)
        
        c_exp_u1, c_exp_u2 = st.columns(2)
        with c_exp_u1:
            if not user_df.empty and REPORTLAB_AVAILABLE:
                staff_pdf = generate_pdf_report(
                    title_text="Field Staff Follow-up Performance Audit",
                    subtitle_text=f"Scope: {scope_dist_txt} | Staff Selected: {user_choice}",
                    sections=[{"title": "Staff Verified Completion Volume & Share", "df": user_df, "summary_stats": global_kpi_summary}],
                )
            else:
                staff_pdf = b""
            st.download_button(
                label="📄 Download Staff PDF",
                data=staff_pdf,
                file_name="TB_Staff_Productivity_Audit.pdf",
                mime="application/pdf",
                width="stretch",
                disabled=(user_df.empty or not REPORTLAB_AVAILABLE),
            )
        with c_exp_u2:
            if not user_df.empty and DOCX_AVAILABLE:
                staff_docx = generate_word_report(
                    title_text="Field Staff Follow-up Performance Audit",
                    subtitle_text=f"Scope: {scope_dist_txt} | Staff Selected: {user_choice}",
                    summary_stats=global_kpi_summary,
                    insights_list=clean_insights,
                    recommendations_list=clean_recs,
                    data_tables=[("Staff Productivity Performance", user_df)],
                )
            else:
                staff_docx = b""
            st.download_button(
                label="📝 Download Staff Word (.docx)",
                data=staff_docx,
                file_name="TB_Staff_Performance.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                width="stretch",
                disabled=(user_df.empty or not DOCX_AVAILABLE),
            )

# --- TAB 5: COMPLETE PATIENT LINE-LIST (100% RECORDS) ---
with tab_patients:
    st.markdown("**Line-Listed High-Risk Patient Directory (Active Filtered Scope)**")
    
    full_export_cols = [c for c in filtered_pt.columns if c not in ["District_Clean", "TU_Clean"]] if not filtered_pt.empty else []
    full_patient_dataset = filtered_pt[full_export_cols] if not filtered_pt.empty else pd.DataFrame()

    st.dataframe(full_patient_dataset, width="stretch", height=360)
    st.caption(f"Showing all **{len(full_patient_dataset):,}** records matching current scope filters.")

    st.markdown("#### 📥 Download Complete Patient Line-List")
    col_p_exp1, col_p_exp2, col_p_exp3 = st.columns(3)

    with col_p_exp1:
        full_csv_payload = full_patient_dataset.to_csv(index=False).encode("utf-8") if not full_patient_dataset.empty else b""
        st.download_button(
            label=f"📋 Download Complete CSV ({len(full_patient_dataset):,} rows)",
            data=full_csv_payload,
            file_name=f"Complete_Patient_LineList_{len(full_patient_dataset)}_records.csv",
            mime="text/csv",
            width="stretch",
            disabled=(len(full_patient_dataset) == 0),
        )

    with col_p_exp2:
        out_pt_excel = io.BytesIO()
        if not full_patient_dataset.empty:
            with pd.ExcelWriter(out_pt_excel, engine="openpyxl") as writer:
                full_patient_dataset.to_excel(writer, sheet_name="Patient Line List", index=False)
            excel_bytes = out_pt_excel.getvalue()
        else:
            excel_bytes = b""
        st.download_button(
            label=f"📊 Download Complete Excel ({len(full_patient_dataset):,} rows)",
            data=excel_bytes,
            file_name=f"Complete_Patient_LineList_{len(full_patient_dataset)}_records.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            width="stretch",
            disabled=(len(full_patient_dataset) == 0),
        )

    with col_p_exp3:
        essential_cols = [c for c in [col_p_id, col_p_dist, col_p_tu, "Follow-up Status", "Follow-up User"] if c and c in full_patient_dataset.columns]
        if not essential_cols and not full_patient_dataset.empty:
            essential_cols = full_patient_dataset.columns[:5].tolist()
        
        pdf_pt_data = full_patient_dataset[essential_cols] if not full_patient_dataset.empty else pd.DataFrame()
        if not pdf_pt_data.empty and REPORTLAB_AVAILABLE:
            line_list_pdf = generate_pdf_report(
                title_text="Patient Line-List Directory Report",
                subtitle_text=f"Total Cohort: {len(pdf_pt_data):,} Records | Scope: {scope_dist_txt}",
                sections=[{"title": "Patient Registry & Verification Details", "df": pdf_pt_data}],
            )
        else:
            line_list_pdf = b""
        st.download_button(
            label=f"📄 Download Patient PDF ({len(pdf_pt_data):,} rows)",
            data=line_list_pdf,
            file_name=f"Patient_LineList_{len(pdf_pt_data)}_records.pdf",
            mime="application/pdf",
            width="stretch",
            disabled=(len(pdf_pt_data) == 0 or not REPORTLAB_AVAILABLE),
        )

st.markdown("---")

# ---------------------------------------------------------
# 15. Consolidated Executive Export Center
# ---------------------------------------------------------
st.markdown("### 📥 Consolidated Executive Export Center")

def build_consolidated_excel(p_df, d_df, t_df, u_df):
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        d_df.to_excel(writer, sheet_name="District Summary", index=False)
        t_df.to_excel(writer, sheet_name="TU Summary", index=False)
        u_df.to_excel(writer, sheet_name="User Performance", index=False)
        p_df.to_excel(writer, sheet_name="Full Patient LineList", index=False)
    return out.getvalue()

c_cen1, c_cen2, c_cen3, c_cen4 = st.columns(4)

with c_cen1:
    if kpi_total_patients > 0 and REPORTLAB_AVAILABLE:
        consolidated_sections = [
            {"title": "1. District Summary & Coverage", "df": district_df, "summary_stats": global_kpi_summary},
            {"title": "2. Tuberculosis Unit (TU) Level Performance", "df": tu_df},
            {"title": "3. Field Staff Follow-up Productivity", "df": user_df},
        ]
        consolidated_pdf = generate_pdf_report(
            title_text="Comprehensive TB Follow-up Executive Analytics Report",
            subtitle_text=f"Scope: {scope_dist_txt} | Focus TU: {tu_choice}",
            sections=consolidated_sections,
        )
    else:
        consolidated_pdf = b""
    st.download_button(
        label="📑 Full Briefing (PDF)",
        data=consolidated_pdf,
        file_name="Comprehensive_TB_Followup_Executive_Report.pdf",
        mime="application/pdf",
        width="stretch",
        disabled=(kpi_total_patients == 0 or not REPORTLAB_AVAILABLE),
    )

with c_cen2:
    if kpi_total_patients > 0 and DOCX_AVAILABLE:
        consolidated_word = generate_word_report(
            title_text="Comprehensive TB Follow-up Executive Analytics Report",
            subtitle_text=f"Scope: {scope_dist_txt} | Focus TU: {tu_choice}",
            summary_stats=global_kpi_summary,
            insights_list=clean_insights,
            recommendations_list=clean_recs,
            data_tables=[
                ("District Performance", district_df),
                ("TU Performance", tu_df),
                ("Staff Output", user_df),
            ],
        )
    else:
        consolidated_word = b""
    st.download_button(
        label="📝 Full Briefing (Word .docx)",
        data=consolidated_word,
        file_name="Comprehensive_TB_Followup_Executive_Report.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        width="stretch",
        disabled=(kpi_total_patients == 0 or not DOCX_AVAILABLE),
    )

with c_cen3:
    if not full_patient_dataset.empty:
        excel_payload = build_consolidated_excel(full_patient_dataset, district_df, tu_df, user_df)
    else:
        excel_payload = b""
    st.download_button(
        label=f"📊 Multi-Sheet Excel ({len(full_patient_dataset):,} rows)",
        data=excel_payload,
        file_name="High_Risk_TB_Full_Analytics_Report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch",
        disabled=(len(full_patient_dataset) == 0),
    )

with c_cen4:
    st.download_button(
        label=f"📋 Complete CSV ({len(full_patient_dataset):,} rows)",
        data=full_csv_payload if not full_patient_dataset.empty else b"",
        file_name="Complete_Filtered_Patient_LineList.csv",
        mime="text/csv",
        width="stretch",
        disabled=(len(full_patient_dataset) == 0),
    )